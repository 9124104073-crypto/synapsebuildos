import os
import random
import docx

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "..", "sample_briefs")
os.makedirs(OUTPUT_DIR, exist_ok=True)

CITIES = ["Chennai", "Bengaluru", "Mumbai", "Delhi-NCR", "Hyderabad", "Pune", "Coimbatore", "Kochi", "Ahmedabad", "Jaipur"]
FACINGS = ["East", "North", "South", "West"]
STYLES = [
    "Modern Minimalist", "Traditional Vastu-Compliant", "Contemporary Glass & Wood",
    "Eco-Friendly Sustainable", "Luxury Mediterranean Villa", "Compact Urban Smart Home",
    "Multi-Generational Family Home", "Tropical Courtyard House"
]

ROOM_CONFIGS = [
    {"bhk": "1 BHK", "rooms": ["Living Room", "Master Bedroom", "Kitchen", "Common Bathroom"], "w": 30, "h": 40, "budget": 35, "floors": 1},
    {"bhk": "2 BHK", "rooms": ["Living Room", "Master Bedroom", "Bedroom 2", "Kitchen", "Dining Room", "Master Bathroom", "Common Bathroom"], "w": 30, "h": 50, "budget": 45, "floors": 1},
    {"bhk": "3 BHK", "rooms": ["Living Room", "Master Bedroom", "Bedroom 2", "Kids Bedroom", "Kitchen", "Dining Room", "Puja Room", "Master Bathroom", "Common Bathroom", "Balcony"], "w": 40, "h": 50, "budget": 75, "floors": 2},
    {"bhk": "4 BHK Villa", "rooms": ["Living Room", "Master Bedroom", "Bedroom 2", "Bedroom 3", "Guest Bedroom", "Kitchen", "Dining Room", "Puja Room", "Home Office", "Master Bathroom", "Attached Bathroom 2", "Common Bathroom", "Balcony", "Terrace"], "w": 50, "h": 60, "budget": 120, "floors": 2},
    {"bhk": "5 BHK Mansion", "rooms": ["Grand Living Room", "Master Suite", "Bedroom 2", "Bedroom 3", "Bedroom 4", "Guest Suite", "Modular Kitchen", "Dining Hall", "Puja Room", "Home Theater", "Home Gym", "Master Bathroom", "Attached Bathroom 2", "Attached Bathroom 3", "Common Bathroom", "Double Balcony", "Terrace Garden"], "w": 60, "h": 80, "budget": 220, "floors": 3}
]

def generate_brief(index):
    config = random.choice(ROOM_CONFIGS)
    city = random.choice(CITIES)
    facing = random.choice(FACINGS)
    style = random.choice(STYLES)
    w = config["w"] + random.choice([-5, 0, 5, 10])
    h = config["h"] + random.choice([-5, 0, 5, 10])
    budget = config["budget"] + random.choice([-10, -5, 0, 10, 25])
    floors = config["floors"]

    title = f"Client Brief #{index:03d} - {config['bhk']} {style} ({w}x{h} ft) in {city}"

    content = f"""SYNAPSE STUDIO - ARCHITECTURAL CLIENT BRIEF
Document Reference: BRIEF-{index:03d}
City / Region: {city}
Style Preference: {style}

1. PROJECT OVERVIEW
-------------------
Client Name: Sample Client #{index:03d}
Project Type: Residential Independent House ({config['bhk']})
Location: {city}, India
Total Stories / Floors: {floors} Floor(s) (G + {floors - 1})

2. PLOT & SITE DETAILS
----------------------
Plot Width: {w} feet
Plot Depth: {h} feet
Plot Total Area: {w * h} sq ft
Plot Facing Direction: {facing}
Target Budget: ₹{budget} Lakhs (₹{budget * 100000:,} INR)

3. ROOM & SPATIAL REQUIREMENTS
------------------------------
Requested Spaces:
"""
    for r in config["rooms"]:
        content += f" - {r}\n"

    content += f"""
4. VASTU & ARCHITECTURAL GUIDELINES
-----------------------------------
- Facing: Main Entrance should face {facing} as per plot orientation.
- Kitchen: Positioned in South-East (Agneya) corner preferred.
- Master Bedroom: Located in South-West (Nairutya) corner on upper floor if applicable.
- Puja Room: North-East (Ishan) corner alignment for maximum positive energy.
- Ventilation: Minimum 15% window-to-floor area ratio for natural cross-ventilation.

5. FINISHES & SPECIAL REQUESTS
------------------------------
- Flooring: Vitrified tiles in living/dining areas, anti-skid ceramic tiles in bathrooms.
- Wall Finish: Premium emulsion paint with accent wall in Living Room.
- Exterior: Weather-proof exterior paint with wooden louvers and balcony glass railings.
- Special Notes: Client requests spacious layout, eco-friendly lighting fixtures, and maximum space efficiency within ₹{budget}L budget.
"""
    return title, content, w, h, budget, facing, config['bhk']

generated_count = 0
brief_manifest = []

for i in range(1, 101):
    title, text, w, h, budget, facing, bhk = generate_brief(i)

    # Save TXT
    txt_filename = f"brief_{i:03d}_{bhk.replace(' ', '_').lower()}_{w}x{h}.txt"
    txt_path = os.path.join(OUTPUT_DIR, txt_filename)
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(text)

    # Save DOCX
    docx_filename = f"brief_{i:03d}_{bhk.replace(' ', '_').lower()}_{w}x{h}.docx"
    docx_path = os.path.join(OUTPUT_DIR, docx_filename)
    doc = docx.Document()
    doc.add_heading(title, level=1)
    for line in text.split("\n"):
        if line.startswith("---") or not line.strip():
            continue
        if line.endswith(":") and len(line) < 40:
            doc.add_heading(line, level=2)
        else:
            doc.add_paragraph(line)
    doc.save(docx_path)

    generated_count += 1
    brief_manifest.append({
        "id": f"brief_{i:03d}",
        "title": title,
        "txt": txt_filename,
        "docx": docx_filename,
        "w": w,
        "h": h,
        "budget": budget,
        "facing": facing,
        "bhk": bhk
    })

print(f"Successfully generated {generated_count} sample client brief documents in {OUTPUT_DIR}")
