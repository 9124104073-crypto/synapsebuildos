import docx
from PIL import Image, ImageDraw

samples = [
    {
        'id': 'compact_2bhk_30x40',
        'title': 'Compact Urban House Brief (2 BHK - 30x40 ft Plot)',
        'filename': 'brief_compact_2bhk_30x40',
        'lines': [
            'CLIENT REQUIREMENT BRIEF - COMPACT URBAN RESIDENCE',
            '',
            'PROJECT SPECIFICATIONS:',
            '- Project Name: City Nest 2BHK',
            '- Location: Madurai, Tamil Nadu',
            '- Site Dimensions: 30 ft x 40 ft (Total Area: 1200 sq ft)',
            '- Road Frontage: North facing road',
            '- Building Type: Single Floor (Ground floor only G+0)',
            '- Bedrooms: 2 BHK (2 bedrooms)',
            '- Bathrooms: 2 bathrooms',
            '- Target Budget: Rs 45 Lakhs',
            '',
            'FAMILY & REQUIREMENTS:',
            '- Family: Family of 4 members',
            '- Rooms required: Living room 14 x 16 ft, Kitchen 9 x 10 ft, Dining 10 x 12 ft, Master Bedroom 12 x 14 ft, Bedroom 2 10 x 12 ft, Pooja room 5 x 6 ft, Car parking 10 x 15 ft',
            '- Finishes: Vitrified tile flooring, warm wall paint'
        ]
    },
    {
        'id': 'modern_3bhk_40x50',
        'title': 'Modern Family Duplex Brief (3 BHK - 40x50 ft Plot)',
        'filename': 'brief_modern_3bhk_40x50',
        'lines': [
            'CLIENT REQUIREMENT BRIEF - MODERN DUPLEX VILLA',
            '',
            'PROJECT SPECIFICATIONS:',
            '- Project Name: Green Valley Residence',
            '- Location: Coimbatore, Tamil Nadu',
            '- Site Dimensions: 40 ft x 50 ft (Total Area: 2000 sq ft)',
            '- Road Frontage: West facing road',
            '- Building Type: G+1 Duplex (2 floors)',
            '- Bedrooms: 3 BHK (3 bedrooms)',
            '- Bathrooms: 3 bathrooms',
            '- Target Budget: Rs 70 Lakhs',
            '',
            'ROOM DETAILS:',
            'Ground Floor:',
            '- Living Room: 16 x 18 ft',
            '- Dining Area: 12 x 14 ft',
            '- Kitchen: 10 x 12 ft',
            '- Master Bedroom 1: 14 x 15 ft with attached bathroom 6 x 8 ft',
            '- Pooja Room: 6 x 8 ft',
            '- Utility: 6 x 8 ft',
            '- Car Parking: 12 x 16 ft',
            'First Floor:',
            '- Master Bedroom 2: 14 x 15 ft with balcony 6 x 12 ft',
            '- Bedroom 3: 12 x 14 ft',
            '- Home Office / Study: 10 x 12 ft',
            '- Bathroom 2 & 3: 6 x 8 ft each'
        ]
    },
    {
        'id': 'luxury_4bhk_50x80',
        'title': 'Luxury Coastal Estate Brief (4 BHK - 50x80 ft Plot)',
        'filename': 'brief_luxury_4bhk_50x80',
        'lines': [
            'CLIENT REQUIREMENT BRIEF - LUXURY COASTAL ESTATE',
            '',
            'PROJECT SPECIFICATIONS:',
            '- Project Name: ECR Luxury Villa',
            '- Location: ECR Chennai, Tamil Nadu',
            '- Site Environment: Coastal site',
            '- Site Dimensions: 50 ft x 80 ft (Total Area: 4000 sq ft)',
            '- Road Frontage: Southeast facing road',
            '- Building Type: G+1 Duplex (2 floors luxury villa)',
            '- Bedrooms: 4 BHK (4 bedrooms)',
            '- Bathrooms: 4 bathrooms',
            '- Target Budget: Rs 1.5 Crores (Rs 150 Lakhs)',
            '',
            'REQUIREMENTS & ROOM SIZES:',
            'Ground Floor:',
            '- Living Room: 20 x 24 ft (Grand double height living)',
            '- Dining Room: 14 x 16 ft',
            '- Kitchen: 12 x 16 ft with pantry',
            '- Master Bedroom 1: 16 x 18 ft (Elderly parents bedroom)',
            '- Bathroom 1: 8 x 10 ft',
            '- Pooja Room: 8 x 10 ft',
            '- Utility: 8 x 10 ft',
            '- Double Car Parking: 16 x 20 ft',
            'First Floor:',
            '- Primary Suite Bedroom 2: 16 x 18 ft with private balcony 10 x 16 ft',
            '- Bedroom 3: 14 x 16 ft',
            '- Bedroom 4: 14 x 16 ft',
            '- Home Office & Library: 12 x 14 ft',
            '- Bathrooms 2, 3, 4: 8 x 10 ft each'
        ]
    },
    {
        'id': 'bungalow_3bhk_45x60',
        'title': 'Single Floor Bungalow Brief (3 BHK - 45x60 ft Plot)',
        'filename': 'brief_bungalow_3bhk_45x60',
        'lines': [
            'CLIENT REQUIREMENT BRIEF - TRADITIONAL BUNGALOW',
            '',
            'PROJECT SPECIFICATIONS:',
            '- Project Name: Royal Heritage Bungalow',
            '- Location: Trichy, Tamil Nadu',
            '- Site Dimensions: 45 ft x 60 ft (Total Area: 2700 sq ft)',
            '- Road Frontage: East facing road',
            '- Building Type: Single storey bungalow (Ground floor only)',
            '- Bedrooms: 3 BHK (3 bedrooms)',
            '- Bathrooms: 3 bathrooms',
            '- Target Budget: Rs 60 Lakhs',
            '',
            'LAYOUT DETAILS:',
            '- Living Room: 18 x 20 ft with traditional Thinnai porch',
            '- Dining Area: 12 x 15 ft',
            '- Kitchen: 11 x 13 ft',
            '- Master Bedroom 1: 14 x 15 ft',
            '- Bedroom 2: 12 x 14 ft',
            '- Bedroom 3: 12 x 14 ft',
            '- Bathrooms: 3 attached baths 6 x 8 ft',
            '- Pooja Room: 6 x 7 ft',
            '- Car Porch: 12 x 18 ft'
        ]
    }
]

for item in samples:
    fname = item['filename']
    base = '../' + fname
    content = '\n'.join(item['lines'])
    
    # TXT
    with open(base + '.txt', 'w', encoding='utf-8') as f:
        f.write(content)
    with open('../web/' + fname + '.txt', 'w', encoding='utf-8') as f:
        f.write(content)
        
    # MD
    with open(base + '.md', 'w', encoding='utf-8') as f:
        f.write('# ' + item['title'] + '\n\n' + content)
        
    # DOCX
    doc = docx.Document()
    doc.add_heading(item['title'], 0)
    for line in item['lines']:
        if line.endswith(':'):
            doc.add_heading(line, level=1)
        elif line.strip():
            doc.add_paragraph(line)
    doc.save(base + '.docx')
    doc.save('../web/' + fname + '.docx')

    # PNG Image
    img = Image.new('RGB', (900, 1200), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    y = 40
    for line in item['lines']:
        if line.endswith(':') or 'CLIENT' in line:
            draw.text((40, y), line, fill=(15, 23, 42))
            y += 30
        else:
            draw.text((40, y), line, fill=(51, 65, 85))
            y += 22
    img.save(base + '.png')
    img.save('../web/' + fname + '.png')

print("All sample brief documents created successfully!")
